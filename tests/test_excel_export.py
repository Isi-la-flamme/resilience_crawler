import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from crawler.database import CrawlDatabase
from crawler.excel_export import export_database_to_excel


class ExcelExportTests(unittest.TestCase):
    def test_exports_pages_images_and_documents(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            database_path = Path(temp_dir) / "crawl.db"
            output_path = Path(temp_dir) / "exports" / "crawl.xlsx"
            start_url = "https://example.org/"
            page_url = "https://example.org/article"
            image_url = "https://cdn.example.org/image.png"

            with CrawlDatabase(database_path) as database:
                database.save_page(start_url, {
                    "url": page_url,
                    "title": "=not a formula",
                    "published_at": "2026-10-01T12:00:00Z",
                    "content": "Collected article content",
                    "html": "<p>Collected article content</p>",
                    "links": [],
                    "images": [image_url],
                }, [])
                database.claim_image(image_url)
                database.finish_image(
                    image_url,
                    Path(temp_dir) / "img" / "image.png",
                )
                database.save_document(start_url, {
                    "url": "https://example.org/report.pdf",
                    "url_hash": CrawlDatabase.page_url_hash(
                        "https://example.org/report.pdf"
                    ),
                    "site_name": "example.org",
                    "local_path": str(Path(temp_dir) / "documents" / "report.pdf"),
                    "filename": "report.pdf",
                    "content_type": "application/pdf",
                    "size_bytes": 1234,
                })

            counts = export_database_to_excel(database_path, output_path)
            workbook = load_workbook(output_path, read_only=False, data_only=False)
            try:
                self.assertEqual(workbook.sheetnames, ["Pages", "Images", "Documents"])
                page_sheet = workbook["Pages"]
                self.assertEqual(page_sheet["D2"].value, "=not a formula")
                self.assertEqual(page_sheet["D2"].data_type, "s")
                self.assertEqual(page_sheet["F2"].value, "Collected article content")
                self.assertFalse(page_sheet["G2"].value)
                self.assertEqual(workbook["Images"]["A2"].value, "cdn.example.org")
                self.assertEqual(workbook["Images"]["D2"].value, "downloaded")
                self.assertEqual(workbook["Documents"]["F2"].value, 1234)
            finally:
                workbook.close()

        self.assertEqual(counts, {"pages": 1, "images": 1, "documents": 1})

    def test_truncates_overlong_content_and_marks_it(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            database_path = Path(temp_dir) / "crawl.db"
            output_path = Path(temp_dir) / "crawl.xlsx"
            with CrawlDatabase(database_path) as database:
                database.save_page("https://example.org/", {
                    "url": "https://example.org/long",
                    "title": "Long page",
                    "content": "x" * 33000,
                    "links": [],
                    "images": [],
                }, [])

            export_database_to_excel(database_path, output_path)
            workbook = load_workbook(output_path, read_only=True)
            try:
                row = next(workbook["Pages"].iter_rows(min_row=2, values_only=True))
            finally:
                workbook.close()

        self.assertEqual(len(row[5]), 32767)
        self.assertTrue(row[5].endswith("..."))
        self.assertTrue(row[6])

    def test_refuses_to_overwrite_database(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            database_path = Path(temp_dir) / "crawl.db"
            connection = sqlite3.connect(database_path)
            connection.close()

            with self.assertRaisesRegex(ValueError, "ne peut pas remplacer"):
                export_database_to_excel(database_path, database_path)


if __name__ == "__main__":
    unittest.main()
