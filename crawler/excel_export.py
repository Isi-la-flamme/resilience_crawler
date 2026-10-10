import logging
import os
import sqlite3
import tempfile
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)
EXCEL_MAX_CELL_LENGTH = 32767
EXPORT_BATCH_SIZE = 1000


def export_database_to_excel(database_path, output_path):
    """Exporte les pages, images et documents SQLite dans un classeur XLSX.

    La base est ouverte en lecture seule et les lignes sont lues par lots pour
    limiter la mémoire. Le classeur temporaire est créé dans le dossier cible,
    puis remplacé atomiquement après sauvegarde complète. La base source ne
    peut pas être remplacée par le classeur.

    Args:
        database_path: Base crawler contenant `pages`, `images`, `documents`.
        output_path: Fichier XLSX à créer ; il peut remplacer un export antérieur.

    Returns:
        Nombre de lignes exportées sous les clés `pages`, `images` et `documents`.

    Raises:
        ValueError: Si la sortie désigne le même fichier que la base source.
        sqlite3.Error, OSError: En cas de problème de lecture ou d'écriture.
    """
    database_path = Path(database_path).resolve()
    output_path = Path(output_path).resolve()
    if database_path == output_path:
        raise ValueError("Le fichier Excel ne peut pas remplacer la base SQLite.")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{output_path.stem}-",
            suffix=".xlsx",
            dir=output_path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)

        workbook = Workbook(write_only=True)
        pages_sheet = _create_sheet(workbook, "Pages", (
            "ID",
            "Site de départ",
            "URL",
            "Titre",
            "Date de publication",
            "Contenu",
            "Contenu tronqué",
            "Date de collecte",
        ))
        images_sheet = _create_sheet(workbook, "Images", (
            "Site",
            "URL",
            "Chemin local",
            "Statut",
            "Erreur",
        ))
        documents_sheet = _create_sheet(workbook, "Documents", (
            "Site de départ",
            "Site",
            "URL",
            "Nom du fichier",
            "Type MIME",
            "Taille (octets)",
            "Chemin local",
            "Date de téléchargement",
        ))

        connection = sqlite3.connect(
            f"{database_path.as_uri()}?mode=ro",
            uri=True,
            timeout=30,
        )
        try:
            counts = {
                "pages": _append_query_rows(
                    connection,
                    pages_sheet,
                    """SELECT id, start_url, url, title, published_at, content, crawled_at
                       FROM pages ORDER BY id""",
                    _page_row,
                ),
                "images": _append_query_rows(
                    connection,
                    images_sheet,
                    """SELECT url, local_path, status, error
                       FROM images ORDER BY url""",
                    _image_row,
                ),
                "documents": _append_query_rows(
                    connection,
                    documents_sheet,
                    """SELECT start_url, site_name, url, filename, content_type,
                              size_bytes, local_path, downloaded_at
                       FROM documents ORDER BY start_url, url""",
                    lambda row: row,
                ),
            }
        finally:
            connection.close()

        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
        logger.info(
            "Excel export saved to %s (pages=%d images=%d documents=%d)",
            output_path,
            counts["pages"],
            counts["images"],
            counts["documents"],
        )
        return counts
    except Exception:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
        raise


def _create_sheet(workbook, title, headers):
    """Crée une feuille filtrable avec ligne d'en-tête figée et mise en forme."""
    sheet = workbook.create_sheet(title)
    sheet.freeze_panes = "A2"
    sheet.append([
        _excel_cell(sheet, value, header=True)
        for value in headers
    ])
    return sheet


def _append_query_rows(connection, sheet, query, row_converter):
    """Ajoute au classeur les résultats SQL par lots et retourne leur nombre."""
    cursor = connection.execute(query)
    count = 0
    while rows := cursor.fetchmany(EXPORT_BATCH_SIZE):
        for row in rows:
            sheet.append([
                _excel_cell(sheet, value)
                for value in row_converter(row)
            ])
            count += 1
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(cursor.description))}{count + 1}"
    )
    return count


def _page_row(row):
    """Convertit une ligne `pages` en colonnes Excel et borne le texte.

    La colonne `Contenu tronqué` indique si le texte a dépassé la limite
    d'Excel ; le contenu est alors coupé en gardant une ellipse finale.
    """
    page_id, start_url, url, title, published_at, content, crawled_at = row
    content_text = "" if content is None else str(content)
    content_truncated = len(content_text) > EXCEL_MAX_CELL_LENGTH
    if content_truncated:
        content_text = content_text[:EXCEL_MAX_CELL_LENGTH - 3] + "..."
    return (
        page_id,
        start_url,
        url,
        title,
        published_at,
        content_text,
        content_truncated,
        crawled_at,
    )


def _image_row(row):
    """Convertit une ligne d'image et dérive le nom du site depuis son URL."""
    url, local_path, status, error = row
    return (
        urlsplit(url).hostname or "",
        url,
        local_path,
        status,
        error,
    )


def _excel_cell(sheet, value, header=False):
    """Nettoie et construit une cellule XLSX, en forçant les chaînes en texte.

    Les caractères illégaux sont supprimés et les textes sont limités à la
    longueur maximale d'une cellule Excel. Les en-têtes reçoivent le style
    visuel du rapport.
    """
    if value is None:
        value = ""
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        if len(value) > EXCEL_MAX_CELL_LENGTH:
            value = value[:EXCEL_MAX_CELL_LENGTH]
    cell = WriteOnlyCell(sheet, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    if header:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    return cell
